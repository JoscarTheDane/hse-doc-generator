# AgentMail key: set AGENTMAIL_API_KEY env var or place key file at $HSE_HOME/.api_key
#!/usr/bin/env python3
"""Extract text from all attachments."""

import os
import sys
sys.path.insert(0, '/tmp/hira_venv/lib/python3.11/site-packages')

SAVE_DIR = "/home/joshua/.hermes/skills/productivity/hse-risk-assessment-generator/scripts/temp_attachments"

files = [
    ("Snake_Island_Wrecks_Removal", "6356-CIN-RFQ-001, Rev. 1.0 - RFQ Wrecks and Objects Removal Snake Island.pdf"),
    ("SNEPCo_Bonga_Mooring", "AM-SNEPCo-Bonga IWS 2023 Mooring Procedure- rev1.docx"),
    ("3D_Photogrammetry_Inspection", "Scope for 3D photogrammetry Inspection_.pdf"),
    ("Flange_Disconnection_Reconnection", "Task Plan for Flange disconnection_recovery and deployment_reconnection.docx"),
]

for prefix, fname in files:
    fpath = os.path.join(SAVE_DIR, f"{prefix}_{fname}")
    if not os.path.exists(fpath):
        print(f"\n{'='*60}")
        print(f"NOT FOUND: {fpath}")
        continue
    
    ext = fname.lower().split('.')[-1]
    print(f"\n{'='*60}")
    print(f"Processing: {fname}")
    print(f"Extension: {ext}")
    
    if ext == 'pdf':
        try:
            from pypdf import PdfReader
            reader = PdfReader(fpath)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            print(text[:8000])
            if len(text) > 8000:
                print(f"\n... (total {len(text)} chars, saving full text)")
                # Save full text
                text_path = os.path.join(SAVE_DIR, f"{prefix}_full.txt")
                with open(text_path, "w") as f:
                    f.write(text)
        except Exception as e:
            print(f"Error: {e}")
    
    elif ext == 'docx':
        from docx import Document
        doc = Document(fpath)
        text = ""
        for para in doc.paragraphs:
            text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                text += row_text + "\n"
        print(text[:8000])
        if len(text) > 8000:
            print(f"\n... (total {len(text)} chars, saving full text)")
            text_path = os.path.join(SAVE_DIR, f"{prefix}_full.txt")
            with open(text_path, "w") as f:
                f.write(text)
